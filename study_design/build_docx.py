#!/usr/bin/env python3
"""
Render BEACON_survey_questionnaire.md into a formatted .docx for client circulation.

Handles the subset of Markdown used in the questionnaire: ATX headings, pipe tables,
bullet and ordered lists, blockquotes, horizontal rules, and inline bold/italic.
"""

import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

SRC = "BEACON_survey_questionnaire.md"
DST = "BEACON_survey_questionnaire.docx"

INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x0B, 0x4F, 0x6C)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)


def add_runs(paragraph, text):
    """Split inline **bold** and *italic* markers into runs."""
    for token in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)


def strip_md(text):
    return re.sub(r"\*\*([^*]+)\*\*", r"\1", re.sub(r"\*([^*]+)\*", r"\1", text))


def split_row(line):
    cells = line.strip().strip("|").split("|")
    return [c.strip() for c in cells]


def is_separator(line):
    return bool(re.fullmatch(r"\|[\s:\-|]+\|", line.strip()))


def main():
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style.font.color.rgb = INK
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.12

    for section in doc.sections:
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)

    lines = open(SRC, encoding="utf-8").read().split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # horizontal rule
        if re.fullmatch(r"-{3,}", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            i += 1
            continue

        # headings
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level, text = len(m.group(1)), strip_md(m.group(2))
            if level == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(text.upper())
                run.bold = True
                run.font.size = Pt(22)
                run.font.color.rgb = ACCENT
                p.paragraph_format.space_after = Pt(2)
            elif level == 2:
                p = doc.add_paragraph()
                run = p.add_run(text)
                run.bold = True
                run.font.size = Pt(14)
                run.font.color.rgb = ACCENT
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(8)
            elif level == 3:
                p = doc.add_paragraph()
                run = p.add_run(text)
                run.bold = True
                run.font.size = Pt(12)
                run.font.color.rgb = ACCENT
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(6)
            else:  # level 4+ : question stems
                p = doc.add_paragraph()
                run = p.add_run(text)
                run.bold = True
                run.font.size = Pt(10.5)
                run.font.color.rgb = INK
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
            i += 1
            continue

        # tables
        if stripped.startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            header = split_row(stripped)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            n_cols = len(header)
            table = doc.add_table(rows=1, cols=n_cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            for c, text in enumerate(header):
                cell = table.rows[0].cells[c]
                cell.text = ""
                run = cell.paragraphs[0].add_run(strip_md(text))
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shading = cell._tc.get_or_add_tcPr()
                from docx.oxml.ns import qn
                from docx.oxml import OxmlElement
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:fill"), "0B4F6C")
                shading.append(shd)
            for r in rows:
                cells = table.add_row().cells
                for c in range(n_cols):
                    text = strip_md(r[c]) if c < len(r) else ""
                    cells[c].text = ""
                    run = cells[c].paragraphs[0].add_run(text)
                    run.font.size = Pt(9)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # blockquote
        if stripped.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(" ".join(strip_md(q) for q in quote if q))
            run.italic = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = MUTED
            continue

        # bullet list
        if stripped.startswith("- "):
            while i < len(lines) and lines[i].strip().startswith("- "):
                p = doc.add_paragraph(style="List Bullet")
                add_runs(p, lines[i].strip()[2:])
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(10)
                i += 1
            continue

        # ordered list
        if re.match(r"^\d+\.\s", stripped):
            while i < len(lines) and re.match(r"^\d+\.\s", lines[i].strip()):
                text = re.sub(r"^\d+\.\s", "", lines[i].strip())
                p = doc.add_paragraph(style="List Number")
                add_runs(p, text)
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(10)
                i += 1
            continue

        # plain paragraph
        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    doc.save(DST)
    print(f"wrote {DST}")


if __name__ == "__main__":
    main()
