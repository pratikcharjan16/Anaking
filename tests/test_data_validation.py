"""Manual validation must never destroy answers or manufacture non-started responses."""
import copy
import csv
import io
import json
import pytest
from models import Study, Respondent, get_db
from test_workspace import CFG, make, approve_publish, login


def start(client, slug, answers=None, complete=False, **extra):
    data = client.post('/api/start', json={'study':slug, **extra}).get_json()
    assert 'session_id' in data, data
    if answers is not None:
        res = client.post('/api/submit' if complete else '/api/save', json={
            'session_id':data['session_id'], 'answers':answers, 'elapsed_seconds':60})
        assert res.status_code == 200, res.get_json()
    return data


def listing(client, slug, **filters):
    res = client.get('/api/workspace/validation', query_string={'study':slug, **filters})
    assert res.status_code == 200, res.get_json()
    return res.get_json()


def decision(client, slug, rid, **extra):
    return client.post('/api/workspace/validation', json={'slug':slug, 'respondent_id':rid,
        'action':'invalidate', 'reason':'speeder', 'notes':'Reviewed manually', 'expected_revision':0, **extra})


def prepare(client):
    slug = make(client)
    approve_publish(client, slug)
    start(client, slug, {'Q1':{'_':1}, 'Q2':{'_':0}}, complete=True)
    rid = listing(client, slug)['rows'][0]['respondent_id']
    return slug, rid


def test_nonstarters_invite_linking_and_lifecycle(app, client):
    slug = make(client)
    approve_publish(client, slug)  # snapshot deliberately has no rid field
    launch = client.post('/api/studio/launch', json={'slug':slug,'recipients':'one@example.org, One\ntwo@example.org, Two'})
    assert launch.status_code == 200
    with app.app_context():
        token = get_db().execute('SELECT token FROM invites WHERE study_slug=? AND email=?', (slug,'one@example.org')).fetchone()[0]
    assert listing(client, slug)['counts']['non_started'] == 2
    d = start(client, slug, embedded={'rid':token})
    rows = listing(client, slug)['rows']
    assert len(rows) == 2
    assert rows[0]['email'] == 'one@example.org' and rows[0]['status'] == 'started'
    assert rows[1]['respondent_id'] is None and rows[1]['status'] == 'non_started'
    rid = rows[0]['respondent_id']
    # Telemetry and empty lists do not constitute an answer; numeric zero does.
    client.post('/api/save', json={'session_id':d['session_id'], 'answers':{'Q1':{'_meta':{'chars':4},'codes':[]}}})
    assert listing(client, slug)['counts']['started'] == 1
    client.post('/api/save', json={'session_id':d['session_id'], 'answers':{'Q2':{'_':0}}})
    assert listing(client, slug)['counts']['in_progress'] == 1
    client.post('/api/submit', json={'session_id':d['session_id'],'answers':{'Q1':{'_':1},'Q2':{'_':0}}})
    assert listing(client, slug)['counts']['complete'] == 1
    # Another launch/reminder record is not another person.
    client.post('/api/studio/launch', json={'slug':slug})
    assert listing(client, slug)['counts']['non_started'] == 1
    start(client, slug, is_test=True, embedded={'rid':token})
    assert listing(client, slug, scope='real')['counts']['total'] == 2
    testrows = listing(client, slug, scope='test')['rows']
    assert len(testrows) == 1 and testrows[0]['email'] == ''
    assert decision(client, slug, None).status_code == 400
    detail = client.get('/api/workspace/validation/response', query_string={'study':slug,'respondent_id':rid}).get_json()
    assert detail['status'] == 'complete' and len(detail['answers']) == 2
    assert '_meta' not in detail['answers'][0]['values']


def test_invalidation_retains_answers_filters_analysis_exports_and_can_restore(app,client):
    slug, rid = prepare(client)
    before = client.get('/api/workspace/analysis?study='+slug).get_json()
    assert decision(client, slug, rid, reason='other', notes='=Manual exclusion for duplicate entry').status_code == 200
    detail_url = f'/api/workspace/validation/response?study={slug}&respondent_id={rid}'
    detail = client.get(detail_url).get_json()
    assert detail['validation']['is_invalid'] and detail['status'] == 'complete'
    assert detail['validation']['reviewer'] == 'local-preview'
    assert detail['history'][0]['action'] == 'invalidate' and detail['answers']
    assert listing(client,slug)['counts']['invalidated'] == 1
    assert client.get('/admin/export.json?study='+slug).get_json() == []
    report = client.get('/api/workspace/analysis?study='+slug).get_json()
    assert report['questions'][0]['base'] == 0 and before['questions'][0]['base'] == 1
    raw = client.get('/admin/export.json?study='+slug+'&include_invalid=1').get_json()
    assert len(raw) == 1 and raw[0]['validation']['is_invalid'] and len(raw[0]['validation_history']) == 1
    assert raw[0]['answers']['Q2']['_'] == '0'
    csvdata = client.get('/admin/export.csv?study='+slug+'&include_invalid=1').data.decode('utf-8-sig')
    exported = list(csv.DictReader(io.StringIO(csvdata)))
    assert exported[0]['validation_status'] == 'invalidated'
    assert exported[0]['validation_notes'].startswith("'=Manual")
    from openpyxl import load_workbook
    xlsx = client.get('/admin/export.xlsx?study='+slug+'&include_invalid=1')
    wb = load_workbook(io.BytesIO(xlsx.data))
    sheet = wb['Responses']
    notes_column = [c.value for c in sheet[1]].index('validation_notes') + 1
    assert sheet.cell(2, notes_column).data_type == 's'
    # Stale and incomplete restoration cannot silently override a review.
    assert decision(client,slug,rid,action='restore',expected_revision=0).status_code == 409
    assert decision(client,slug,rid,action='restore',expected_revision=1,notes='').status_code == 400
    assert decision(client,slug,rid,action='restore',expected_revision=1,notes='Verified answer quality').status_code == 200
    assert len(client.get('/admin/export.json?study='+slug).get_json()) == 1
    restored = client.get(detail_url).get_json()
    assert not restored['validation']['is_invalid'] and len(restored['history']) == 2
    assert [h['action'] for h in restored['history']] == ['restore','invalidate']
    assert decision(client,slug,rid,action='restore',expected_revision=2).status_code == 409


@pytest.mark.parametrize('patch', [
    {'reason':'unknown'}, {'reason':[]}, {'reason':'other','notes':'  '},
    {'notes':'x'*2001}, {'notes':{}}, {'action':'delete'}, {'expected_revision':-1},
    {'expected_revision':'0'}, {'expected_revision':True}, {'respondent_id':True}, {'respondent_id':'R001'}])
def test_validation_rejects_bad_decisions(client,patch):
    slug,rid = prepare(client)
    assert decision(client,slug,rid,**patch).status_code == 400
    assert listing(client,slug)['counts']['invalidated'] == 0


@pytest.mark.parametrize('reason', ['speeder','gibberish','non_contextual','flat_lining','other'])
def test_all_prelisted_reasons(client,reason):
    slug,rid = prepare(client)
    assert decision(client,slug,rid,reason=reason).status_code == 200
    assert listing(client,slug)['rows'][0]['validation']['reason'] == reason


def test_pagination_search_scope_and_filtered_list_export(app,client):
    slug,rid = prepare(client)
    with app.app_context():
        study = Study.get(slug)
        for _ in range(55): Respondent.create(study, False)
    data = listing(client,slug,page_size=10,page=2)
    assert len(data['rows']) == 10 and data['pages'] == 6 and data['total'] == 56
    assert listing(client,slug,q='R001')['rows'][0]['respondent_id'] == rid
    assert listing(client,slug,status='started')['total'] == 55
    decision(client,slug,rid)
    assert listing(client,slug,validity='invalidated')['total'] == 1
    assert listing(client,slug,validity='included')['total'] == 55
    rows = list(csv.DictReader(io.StringIO(client.get('/api/workspace/validation',query_string={
        'study':slug,'validity':'invalidated','format':'csv'}).data.decode('utf-8-sig'))))
    assert len(rows) == 1 and rows[0]['response_id'] == str(rid)
    assert client.get('/api/workspace/validation?study='+slug+'&page_size=1000').status_code == 400
    assert client.get('/api/workspace/validation?study='+slug+'&status=bogus').status_code == 400


@pytest.mark.parametrize('role, expected', [('owner',200),('editor',200),('analyst',200),('reviewer',403),('viewer',403)])
def test_validation_roles_csrf_and_study_boundaries(app,client,role,expected):
    slug,rid = prepare(client)
    other = make(client,'another-study')
    headers = login(app,client,'staff-'+role,role,[slug])
    assert client.get('/api/workspace/validation?study='+slug).status_code == expected
    assert client.get(f'/api/workspace/validation/response?study={slug}&respondent_id={rid}').status_code == expected
    assert client.get('/api/workspace/validation?study='+other).status_code == 403
    body = {'slug':slug,'respondent_id':rid,'action':'invalidate','reason':'gibberish','expected_revision':0}
    assert client.post('/api/workspace/validation',json=body).status_code == 403
    assert client.post('/api/workspace/validation',json=body,headers=headers).status_code == expected
    if expected == 200:
        assert listing(client,slug)['rows'][0]['validation']['reviewer'] == 'staff-'+role


def test_response_scope_and_cleanup(app,client):
    slug,rid = prepare(client)
    other = make(client,'other-validation-study')
    assert decision(client,other,rid).status_code == 404
    assert client.get(f'/api/workspace/validation/response?study={other}&respondent_id={rid}').status_code == 404
    decision(client,slug,rid)
    assert client.post('/api/studio/delete',json={'slug':slug}).status_code == 200
    with app.app_context():
        assert get_db().execute('SELECT count(*) FROM response_validations').fetchone()[0] == 0
        assert get_db().execute('SELECT count(*) FROM response_validation_events').fetchone()[0] == 0


def test_in_progress_invalidation_persists_when_participant_finishes_and_reset_cleans_history(app,client):
    slug=make(client);approve_publish(client,slug)
    d=start(client,slug,{'Q2':{'_':0}})
    rid=listing(client,slug)['rows'][0]['respondent_id']
    assert decision(client,slug,rid,reason='flat_lining').status_code == 200
    assert client.post('/api/submit',json={'session_id':d['session_id'],'answers':{'Q1':{'_':2},'Q2':{'_':5}}}).status_code == 200
    row=listing(client,slug)['rows'][0]
    assert row['status']=='complete' and row['validation']['is_invalid']
    assert client.get('/admin/export.json?study='+slug).get_json()==[]
    with app.app_context():
        assert Respondent.reset(Study.get(slug),'test') == 0
        assert get_db().execute('SELECT count(*) FROM response_validations').fetchone()[0]==1
        assert Respondent.reset(Study.get(slug),'real') == 1
        assert get_db().execute('SELECT count(*) FROM response_validation_events').fetchone()[0]==0


def test_revalidation_history_stale_decision_and_idempotent_migration(app,client):
    from models import init_db
    slug,rid=prepare(client)
    assert decision(client,slug,rid).status_code==200
    assert decision(client,slug,rid,reason='other',notes='New reason').status_code==409
    assert decision(client,slug,rid,reason='non_contextual',expected_revision=1).status_code==200
    with app.app_context():
        init_db()  # additive schema must preserve existing validation decisions
    detail=client.get(f'/api/workspace/validation/response?study={slug}&respondent_id={rid}').get_json()
    assert detail['validation']['revision']==2 and len(detail['history'])==2
    assert detail['validation']['reason']=='non_contextual'
    assert client.post('/api/workspace/validation',json=['invalid shape']).status_code==400
