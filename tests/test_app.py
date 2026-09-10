"""
In-process tests for the Flask platform (no running server needed).

    python3 -m pytest
"""

import base64
import io

from openpyxl import load_workbook

PILOT_CFG = {
    "sections": [{"id": "S1", "title": "Intro"}],
    "questions": [
        {"id": "Q1", "section": "S1", "type": "single_select", "stem": "Pick one",
         "options": [{"code": 1, "label": "A"}, {"code": 2, "label": "B"}]},
        {"id": "Q2", "section": "S1", "type": "open_text", "stem": "Why?"},
    ],
    "qc": {"min_seconds": 10, "verbatim_qs": ["Q2"]},
}


def _start(client, study="beacon", is_test=True):
    r = client.post("/api/start", json={"study": study, "is_test": is_test})
    assert r.status_code == 200, r.get_json()
    return r.get_json()


# ---------------------------------------------------------------- pages & assets
def test_health(client):
    assert client.get("/healthz").get_json()["ok"] is True


def test_survey_pages_inject_study_slug(client):
    for path in ("/", "/test", "/s/beacon", "/s/beacon/test"):
        r = client.get(path)
        assert r.status_code == 200
        assert b'window.STUDY={slug:"beacon"}' in r.data
        assert b"/static/js/survey.js" in r.data
    assert client.get("/s/does-not-exist").status_code == 404


def test_static_assets_served(client):
    assert client.get("/static/js/survey.js").status_code == 200
    assert client.get("/static/css/survey.css").mimetype == "text/css"
    assert client.get("/survey.js").status_code == 404   # old flat paths are gone
    assert client.get("/audio/welcome.mp3").mimetype == "audio/mpeg"
    r = client.get("/audio/welcome.mp3", headers={"Range": "bytes=0-9"})
    assert r.status_code == 206 and len(r.data) == 10
    assert client.get("/audio/app.py").status_code == 404


def test_api_errors_are_json(client):
    r = client.get("/api/spec/nope")
    assert r.status_code == 404 and r.get_json()["error"] == "unknown study"
    assert client.get("/api/does-not-exist").status_code == 404
    assert client.get("/api/does-not-exist").is_json


# ---------------------------------------------------------------- auth
def test_admin_and_studio_need_token(client, token):
    for path in ("/admin", "/studio", "/api/admin/data", "/api/studio/list",
                 "/admin/export.xlsx", "/admin/export.csv"):
        assert client.get(path).status_code == 403, path
        assert client.get(path, query_string={"token": "wrong"}).status_code == 403, path
        assert client.get(path, query_string={"token": token}).status_code == 200, path
    assert client.post("/admin/reset").status_code == 403
    assert client.post("/api/studio/save", json={}).status_code == 403


# ---------------------------------------------------------------- respondent flow
def test_beacon_seeded_and_spec_shape(client):
    spec = client.get("/api/spec").get_json()
    assert len(spec["questions"]) >= 20
    assert spec["conjoint"]["n_tasks"] == 9
    assert isinstance(spec["conjoint"]["tasks"], dict)   # engine addresses tasks by key
    assert spec["terminate_text"]


def test_start_assigns_codes_and_design(client):
    a = _start(client, is_test=True)
    b = _start(client, is_test=False)
    c = _start(client, is_test=True)
    assert (a["respondent_code"], b["respondent_code"], c["respondent_code"]) == \
        ("T001", "R001", "T002")
    assert sorted(b["task_order"]) == list(range(1, 10))
    assert b["from_prebuilt_map"] is True   # R001 comes from respondent_task_map.csv
    assert len(b["alt_positions"]) == 9


def test_save_submit_progress_roundtrip(client, token):
    s = _start(client)
    sid = s["session_id"]
    r = client.post("/api/save", json={"session_id": sid, "elapsed_seconds": 30,
                                       "answers": {"Q12": {"_": "5"},
                                                   "Q4": {"codes": [1, 3]}}})
    assert r.get_json() == {"ok": True, "status": "in_progress"}
    p = client.get("/api/progress", query_string={"sid": sid}).get_json()
    assert p["exists"] and p["answers"]["Q12"]["_"] == "5"
    assert p["answers"]["Q4"]["codes"] == [1, 3]

    r = client.post("/api/submit", json={"session_id": sid, "elapsed_seconds": 60})
    body = r.get_json()
    assert body["ok"] and body["respondent_code"] == "T001"
    assert "speeder" in body["flags"]           # 60 s << 480 s minimum

    data = client.get("/api/admin/data", query_string={"token": token}).get_json()
    assert data["counts"] == {"total": 1, "complete": 1, "screened_out": 0, "in_progress": 0}
    assert client.get("/api/progress", query_string={"sid": "nope"}).get_json() == \
        {"exists": False}


def test_unknown_session_rejected(client):
    assert client.post("/api/save", json={"session_id": "x"}).status_code == 400
    assert client.post("/api/submit", json={"session_id": "x"}).status_code == 400
    assert client.post("/api/save", data="not json",
                       content_type="application/json").status_code == 400


def test_voice_upload(client, token, app):
    s = _start(client)
    blob = base64.b64encode(b"\x1aE\xdf\xa3fake-webm").decode()
    r = client.post("/api/voice", json={"session_id": s["session_id"], "qid": "Q8b",
                                        "ext": "webm", "data": blob})
    assert r.get_json()["file"] == "beacon__T001_Q8b.webm"
    r = client.get("/admin/voice/beacon__T001_Q8b.webm", query_string={"token": token})
    assert r.status_code == 200 and r.mimetype == "audio/webm"
    bad = client.post("/api/voice", json={"session_id": s["session_id"], "qid": "../x",
                                          "data": blob})
    assert bad.status_code == 400
    bad = client.post("/api/voice", json={"session_id": s["session_id"], "qid": "Q1",
                                          "data": "%%%not-base64%%%"})
    assert bad.status_code == 400


# ---------------------------------------------------------------- studio lifecycle
def test_studio_create_launch_respond_delete(client, token):
    q = {"token": token}
    r = client.post("/api/studio/save", query_string=q,
                    json={"title": "Pilot Study", "cfg": PILOT_CFG})
    assert r.get_json() == {"ok": True, "slug": "pilot-study"}

    # drafts are hidden from the public but visible with the preview token
    assert client.get("/s/pilot-study").status_code == 403
    assert client.get("/s/pilot-study", query_string={"preview": token}).status_code == 200
    r = client.post("/api/start", json={"study": "pilot-study"})
    assert r.status_code == 403 and r.get_json()["error"] == "study not live"

    r = client.post("/api/studio/status", query_string=q,
                    json={"slug": "pilot-study", "status": "live"})
    assert r.get_json()["ok"]
    s = _start(client, "pilot-study")
    assert s["task_order"] == []   # no conjoint in this study
    client.post("/api/save", json={"session_id": s["session_id"], "elapsed_seconds": 5,
                                   "answers": {"Q1": {"_": "2"},
                                               "Q2": {"_": "qwertyuiop asdfghjkl zxcvbnm"}}})
    done = client.post("/api/submit", json={"session_id": s["session_id"],
                                            "elapsed_seconds": 5}).get_json()
    assert set(done["flags"]) == {"speeder", "gibberish_verbatim_Q2"}

    # data is isolated per study
    listing = {x["slug"]: x for x in client.get("/api/studio/list",
                                                 query_string=q).get_json()}
    assert listing["pilot-study"]["complete"] == 1
    assert listing["beacon"]["started"] == 0

    csv_txt = client.get("/admin/export.csv",
                         query_string={"token": token, "study": "pilot-study"}).data.decode()
    header, row = csv_txt.splitlines()[:2]
    rec = dict(zip(header.split(","), row.split(",")))
    assert rec["Q1"] == "2" and rec["Q1_text"] == "B" and rec["is_test"] == "test"
    assert "gibberish_verbatim_Q2" in rec["qc_flags"]

    assert client.post("/api/studio/delete", query_string=q,
                       json={"slug": "beacon"}).status_code == 400
    assert client.post("/api/studio/delete", query_string=q,
                       json={"slug": "pilot-study"}).get_json()["ok"]
    assert client.get("/s/pilot-study").status_code == 404


def test_studio_validation(client, token):
    q = {"token": token}
    bad = dict(PILOT_CFG, questions=[{"id": "Q1", "section": "S1", "type": "open_text",
                                      "stem": ""}] * 2)
    r = client.post("/api/studio/save", query_string=q, json={"title": "x", "cfg": bad})
    assert r.status_code == 400 and "duplicate" in r.get_json()["error"]
    bad = dict(PILOT_CFG, questions=[{"id": "Q1", "section": "NOPE", "type": "open_text",
                                      "stem": ""}])
    r = client.post("/api/studio/save", query_string=q, json={"title": "x", "cfg": bad})
    assert r.status_code == 400 and "unknown section" in r.get_json()["error"]
    r = client.post("/api/studio/status", query_string=q,
                    json={"slug": "beacon", "status": "bogus"})
    assert r.status_code == 400


def test_make_conjoint_is_balanced(client, token):
    attrs = [{"id": "price", "levels": ["$1", "$2", "$3"], "higher_is_bad": True},
             {"id": "efficacy", "levels": ["low", "mid", "high"]}]
    r = client.post("/api/studio/make_conjoint", query_string={"token": token},
                    json={"attributes": attrs, "n_tasks": 9, "seed": 3})
    d = r.get_json()
    assert len(d["tasks"]) == 9 and all(len(t) == 3 for t in d["tasks"])
    # exact level balance: 9 tasks x 3 alts = 27 profiles, 3 levels -> 9 each
    for attr in ("price", "efficacy"):
        assert list(d["balance"][attr].values()) == [9, 9, 9]
    # no alternative fully dominates another within a task
    for task in d["tasks"]:
        for x in task:
            for y in task:
                if x is y:
                    continue
                gx = (2 - x["price"], x["efficacy"])
                gy = (2 - y["price"], y["efficacy"])
                assert not (gx[0] >= gy[0] and gx[1] >= gy[1] and gx != gy), task
    r = client.post("/api/studio/make_conjoint", query_string={"token": token},
                    json={"attributes": [{"id": "x"}]})
    assert r.status_code == 400


# ---------------------------------------------------------------- exports & reset
def test_xlsx_export_and_reset(client, token):
    _start(client, is_test=True)
    real = _start(client, is_test=False)
    client.post("/api/submit", json={"session_id": real["session_id"], "elapsed_seconds": 900})

    r = client.get("/admin/export.xlsx", query_string={"token": token, "scope": "all"})
    assert r.status_code == 200 and r.data[:2] == b"PK"
    assert "spreadsheetml" in r.mimetype
    assert r.headers["X-Export-Backend"] in ("openpyxl", "stdlib")
    wb = load_workbook(io.BytesIO(r.data))
    assert {"Field summary", "Responses", "Conjoint long", "Screen-outs", "QC flags",
            "Data dictionary"} <= set(wb.sheetnames)
    assert wb["Responses"].max_row - 1 == 2

    r = client.get("/admin/export.json", query_string={"token": token, "scope": "real"})
    body = r.get_json()
    assert [x["respondent_code"] for x in body] == ["R001"]
    assert "flat" not in body[0]

    r = client.post("/admin/reset", query_string={"token": token, "scope": "test"})
    assert r.get_json()["deleted_respondents"] == 1
    r = client.post("/admin/reset", query_string={"token": token, "scope": "bogus"})
    assert r.status_code == 400
    data = client.get("/api/admin/data", query_string={"token": token}).get_json()
    assert data["counts"]["total"] == 1
