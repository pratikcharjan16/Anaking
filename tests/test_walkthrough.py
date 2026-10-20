"""
Tests for the Walk Through question type.

    python3 -m pytest tests/test_walkthrough.py

Covers turning a PowerPoint deck into narrated steps (core.walkthrough), the shapes the
respondent side reads (static/js/survey.js walkStepsOf / walkQuestionsOf, mirrored here by
core.walkthrough.walk_steps / walk_questions), the flatten + data-document branches in
core.reporting, and the Studio upload endpoint (/api/studio/walkthrough/ppt).
"""

import io

import pytest
from pptx import Presentation
from pptx.util import Inches

from core.reporting import flatten
from core.walkthrough import DECK_MAX_BYTES, DeckError, steps_from_pptx, walk_questions, walk_steps


# ---------------------------------------------------------------- deck fixtures
def deck_bytes(slides):
    """A .pptx in memory: [(title, body, notes)] per slide."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for title, body, notes in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


DECK = deck_bytes([
    ("Mechanism of action", "It blocks the target before the tumour can grow.", "Say this slowly."),
    ("Headline efficacy", "45% versus 25% progression-free survival at 12 months.", ""),
    ("Safety", "Grade 3+ adverse events in about 30% of patients.", "Pause here."),
])


# ---------------------------------------------------------------- steps_from_pptx
def test_one_step_per_slide_with_title_and_narration():
    steps = steps_from_pptx(DECK)
    assert [s["label"] for s in steps] == [
        "Mechanism of action", "Headline efficacy", "Safety"]
    assert steps[0]["narration"] == (
        "It blocks the target before the tumour can grow. Say this slowly.")
    # the speaker notes are part of what gets spoken, in slide order
    assert steps[1]["narration"] == "45% versus 25% progression-free survival at 12 months."
    assert steps[2]["narration"].endswith("Pause here.")


def test_a_blank_slide_is_skipped_rather_than_shown_silent():
    data = deck_bytes([("Only a title", "", ""), ("", "", "")])
    steps = steps_from_pptx(data)
    assert [s["label"] for s in steps] == ["Only a title"]
    assert steps[0]["narration"] == ""


def test_a_deck_with_no_text_at_all_is_refused():
    with pytest.raises(DeckError):
        steps_from_pptx(deck_bytes([("", "", ""), ("", "", "")]))


@pytest.mark.parametrize("blob", [b"", b"not a deck at all", b"PK\x03\x04garbage"])
def test_unreadable_files_are_refused_with_a_reason(blob):
    with pytest.raises(DeckError):
        steps_from_pptx(blob)


def test_a_huge_file_is_refused_before_it_is_parsed():
    with pytest.raises(DeckError):
        steps_from_pptx(b"x" * (DECK_MAX_BYTES + 1))


def test_the_step_count_is_capped():
    data = deck_bytes([("Slide %d" % i, "body %d" % i, "") for i in range(80)])
    assert len(steps_from_pptx(data, max_steps=5)) == 5


# ---------------------------------------------------------------- question shapes
WT_Q = {
    "id": "WT1",
    "section": "S1",
    "type": "walk_through",
    "stem": "A short walkthrough",
    "required": False,
    "walk_tts": True,
    "walk_steps": [
        {"id": "WT1_W1", "label": "Step one", "narration": "Listen to this.",
         "art": "mechanism", "duration": "no-limit", "allow_replay": True, "media": None},
        {"id": "WT1_W2", "label": "Step two", "narration": "And then this.",
         "media": {"src": "/media/s/m_1.png", "kind": "image", "external": False}},
        {"id": "WT1_W3", "label": "Step three", "narration": "Finally this.",
         "src": "/narration/s/clip_deadbeef.mp3", "seconds": 12.5},
    ],
    "walk_questions": [
        {"id": "WT1_WQ1", "type": "open_text", "stem": "What stood out?", "required": False},
    ],
}


def test_steps_are_normalised_for_the_player():
    steps = walk_steps(WT_Q)
    assert [s["label"] for s in steps] == ["Step one", "Step two", "Step three"]
    # the explainer speaks `caption`, so the narration has to land there
    assert [s["caption"] for s in steps] == ["Listen to this.", "And then this.", "Finally this."]
    # a step with no artwork falls back to the generic scene
    assert steps[1]["art"] == "generic"
    # a step with an attachment carries it through untouched
    assert steps[1]["media"]["src"] == "/media/s/m_1.png"
    # a recorded clip keeps its own length
    assert steps[2]["src"].endswith(".mp3") and steps[2]["seconds"] == 12.5


def test_a_step_missing_its_keys_still_renders():
    steps = walk_steps({"id": "WT", "walk_steps": [{"narration": "hi"}]})
    assert steps[0]["label"] == "Step 1"
    assert steps[0]["caption"] == "hi"
    assert steps[0]["art"] == "generic"


def test_an_older_saved_question_still_reads():
    legacy = {"id": "WT", "walkthrough_steps": [{"label": "Old", "narration": "still works"}],
              "walkthrough_questions": [{"id": "WT_WQ1", "type": "open_text"}]}
    assert [s["label"] for s in walk_steps(legacy)] == ["Old"]
    assert [q["id"] for q in walk_questions(legacy)] == ["WT_WQ1"]


def test_no_steps_means_nothing_to_play():
    assert walk_steps({"id": "WT"}) == []
    assert walk_questions({"id": "WT"}) == []


# ---------------------------------------------------------------- export
RESPONDENT = {
    "respondent_code": "T001", "is_test": 1, "status": "complete",
    "screen_out_at": None, "screen_out_reason": None, "started_at": "2026-01-01 10:00:00",
    "completed_at": "2026-01-01 10:06:00", "elapsed_seconds": 360, "embedded": "{}",
    "language": "en-US",
}

ANSWERS = {
    "WT1_WQ1": {"_": "the dosing schedule"},
    "WT1_WQ2": {"_": "2"},
}


def test_the_questions_after_a_walkthrough_are_exported():
    q = dict(WT_Q, walk_questions=WT_Q["walk_questions"] + [
        {"id": "WT1_WQ2", "type": "single_select",
         "options": [{"code": 1, "label": "Yes"}, {"code": 2, "label": "No"}]}])
    out = flatten(RESPONDENT, ANSWERS, {"questions": [q]})
    assert out["WT1_WQ1"] == "the dosing schedule"
    assert out["WT1_WQ2"] == "2"


def test_a_walkthrough_stores_nothing_of_its_own():
    # the walkthrough question has no id of its own in the answers - only its questions do
    out = flatten(RESPONDENT, ANSWERS, {"questions": [WT_Q]})
    assert "WT1" not in out
    assert out["WT1_WQ1"] == "the dosing schedule"


# ---------------------------------------------------------------- endpoint
def _mk(client, slug, cfg):
    r = client.post("/api/studio/save", json={"title": slug, "cfg": cfg})
    assert r.status_code == 200, r.get_data(as_text=True)
    return r.get_json()["slug"]


WT_CFG = {
    "sections": [{"id": "S1", "title": "Main"}],
    "questions": [WT_Q],
}


def test_a_deck_upload_becomes_narrated_steps(client):
    slug = _mk(client, "wt-ppt", WT_CFG)
    r = client.post("/api/studio/walkthrough/ppt?study=" + slug,
                    data={"file": (io.BytesIO(DECK), "deck.pptx")},
                    content_type="multipart/form-data")
    assert r.status_code == 200, r.get_data(as_text=True)
    body = r.get_json()
    assert body["slides"] == 3
    assert [s["label"] for s in body["steps"]] == [
        "Mechanism of action", "Headline efficacy", "Safety"]
    assert "blocks the target" in body["steps"][0]["narration"]


def test_an_unknown_study_is_refused(client):
    r = client.post("/api/studio/walkthrough/ppt?study=nope",
                    data={"file": (io.BytesIO(DECK), "deck.pptx")},
                    content_type="multipart/form-data")
    assert r.status_code == 404


def test_a_non_deck_upload_is_refused(client):
    slug = _mk(client, "wt-bad", WT_CFG)
    r = client.post("/api/studio/walkthrough/ppt?study=" + slug,
                    data={"file": (io.BytesIO(b"not a deck"), "notes.txt")},
                    content_type="multipart/form-data")
    assert r.status_code == 400
    assert ".pptx" in r.get_json()["error"]


def test_a_corrupt_deck_is_refused_with_the_reason(client):
    slug = _mk(client, "wt-corrupt", WT_CFG)
    r = client.post("/api/studio/walkthrough/ppt?study=" + slug,
                    data={"file": (io.BytesIO(b"garbage"), "deck.pptx")},
                    content_type="multipart/form-data")
    assert r.status_code == 400
    assert r.get_json()["error"]
