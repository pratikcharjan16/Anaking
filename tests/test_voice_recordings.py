import base64
import csv
import io
import json

from openpyxl import load_workbook


VOICE_CFG = {
    "sections": [{"id": "S1", "title": "Answers"}],
    "questions": [
        {"id": "Q1", "section": "S1", "type": "open_text", "stem": "Tell us why"},
        {"id": "Q2", "section": "S1", "type": "text_grid", "stem": "Tell us more",
         "rows": [{"code": "r1", "label": "Access"}],
         "cols": [{"code": "c1", "label": "Today"}]},
    ],
}


def test_voice_audio_and_transcripts_export_in_question_section(client):
    created = client.post("/api/studio/save", json={"title": "Voice Export Test",
                                                    "cfg": VOICE_CFG}).get_json()
    slug = created["slug"]
    session = client.post("/api/start", json={"study": slug, "is_test": True}).get_json()
    audio = base64.b64encode(b"test-audio-bytes").decode()

    main_upload = client.post("/api/voice", json={
        "session_id": session["session_id"], "qid": "Q1", "data": audio, "ext": "webm",
    }).get_json()
    grid_upload = client.post("/api/voice", json={
        "session_id": session["session_id"], "qid": "Q2", "item": "r1_c1",
        "data": audio, "ext": "webm",
    }).get_json()
    assert main_upload["ok"] and main_upload["file"].endswith("_Q1.webm")
    assert grid_upload["ok"] and grid_upload["file"].endswith("_Q2_r1_c1.webm")
    assert client.get("/admin/voice/" + grid_upload["file"]).data == b"test-audio-bytes"
    assert client.post("/api/voice", json={
        "session_id": session["session_id"], "qid": "Q2", "item": "../unsafe",
        "data": audio,
    }).status_code == 400

    answers = {
        "Q1": {"_": "Spoken main answer", "transcript": "Transcript for main answer",
               "voice": main_upload["file"]},
        "Q2": {"r1_c1": "Spoken grid answer", "r1_c1_transcript": "Transcript for grid",
               "r1_c1_voice": grid_upload["file"]},
    }
    saved = client.post("/api/save", json={"session_id": session["session_id"],
                                           "answers": answers})
    assert saved.status_code == 200 and saved.get_json()["ok"]
    submitted = client.post("/api/submit", json={"session_id": session["session_id"],
                                                  "elapsed_seconds": 300})
    assert submitted.status_code == 200

    csv_response = client.get("/admin/export.csv", query_string={"study": slug, "scope": "test"})
    export_row = next(csv.DictReader(io.StringIO(csv_response.data.decode())))
    assert export_row["Q1_transcript"] == "Transcript for main answer"
    assert export_row["Q2_r1_c1_transcript"] == "Transcript for grid"
    assert export_row["Q2_r1_c1_voice_file"].endswith("_Q2_r1_c1.webm")

    workbook = load_workbook(io.BytesIO(client.get(
        "/admin/export.xlsx", query_string={"study": slug, "scope": "test"}).data))
    assert "Voice transcripts" in workbook.sheetnames
    sheet = list(workbook["Voice transcripts"].iter_rows(values_only=True))
    header = sheet[0]
    transcript_col = header.index("transcript")
    question_col = header.index("question_id")
    assert {(row[question_col], row[transcript_col]) for row in sheet[1:]} == {
        ("Q1", "Transcript for main answer"), ("Q2", "Transcript for grid")}

    raw = json.loads(client.get("/admin/export.json",
                                query_string={"study": slug, "scope": "test"}).data)
    assert len(raw) == 1
    assert [v["question_id"] for v in raw[0]["voice_transcripts"]] == ["Q1", "Q2"]


def test_landing_hides_admin_portal_link_but_shows_admin_login(client):
    response = client.get("/")
    assert response.status_code == 200
    assert b"admin-portal" not in response.data.lower()
    assert b"/admin-login" in response.data
