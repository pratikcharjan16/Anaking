#!/usr/bin/env python3
"""
PROJECT BEACON - Excel (.xlsx) export builder.

Sheets are defined once as (name, headers, rows, widths) and rendered by whichever backend is
available: openpyxl when installed, otherwise a minimal OOXML writer built on the standard
library's zipfile. Either way the caller gets a real, valid .xlsx - never a renamed CSV.
"""

from __future__ import annotations

import io
import zipfile
from xml.sax.saxutils import escape

HEADER_FILL = "0B4F6C"
HEADER_FONT = "FFFFFF"


# ======================================================================================
# STANDARD-LIBRARY FALLBACK WRITER
# ======================================================================================
class MiniXlsx:
    """Minimal but spec-valid multi-sheet .xlsx writer using only the standard library."""

    def __init__(self):
        self.sheets: list[tuple[str, list[list]]] = []

    def add_sheet(self, name: str, rows: list[list]) -> None:
        safe = name[:31]
        for ch in '[]:*?/\\':
            safe = safe.replace(ch, "_")
        self.sheets.append((safe, rows))

    @staticmethod
    def _col_letter(n: int) -> str:
        s = ""
        while n > 0:
            n, r = divmod(n - 1, 26)
            s = chr(65 + r) + s
        return s

    @classmethod
    def _cell(cls, ref: str, value, bold: bool = False) -> str:
        style = ' s="1"' if bold else ""
        if value is None or value == "":
            return f'<c r="{ref}"{style}/>'
        if isinstance(value, bool):
            return f'<c r="{ref}"{style} t="b"><v>{int(value)}</v></c>'
        if isinstance(value, (int, float)):
            return f'<c r="{ref}"{style}><v>{value}</v></c>'
        text = escape(str(value))
        return f'<c r="{ref}"{style} t="inlineStr"><is><t xml:space="preserve">{text}</t></is></c>'

    def _sheet_xml(self, rows: list[list]) -> str:
        out = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
               '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">',
               '<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" '
               'activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>',
               '<sheetData>']
        for r_i, row in enumerate(rows, start=1):
            cells = "".join(self._cell(f"{self._col_letter(c_i)}{r_i}", v, bold=(r_i == 1))
                            for c_i, v in enumerate(row, start=1))
            out.append(f'<row r="{r_i}">{cells}</row>')
        out.append("</sheetData></worksheet>")
        return "".join(out)

    def save_bytes(self) -> bytes:
        n = len(self.sheets)
        content_types = [
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">',
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>',
            '<Default Extension="xml" ContentType="application/xml"/>',
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-'
            'officedocument.spreadsheetml.sheet.main+xml"/>',
            '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-'
            'officedocument.spreadsheetml.styles+xml"/>',
        ]
        for i in range(1, n + 1):
            content_types.append(
                f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/'
                f'vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>')
        content_types.append("</Types>")

        workbook = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
                    '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">',
                    "<sheets>"]
        for i, (name, _) in enumerate(self.sheets, start=1):
            workbook.append(f'<sheet name="{escape(name)}" sheetId="{i}" r:id="rId{i}"/>')
        workbook.append("</sheets></workbook>")

        wb_rels = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">']
        for i in range(1, n + 1):
            wb_rels.append(
                f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/'
                f'officeDocument/2006/relationships/worksheet" Target="worksheets/sheet{i}.xml"/>')
        wb_rels.append(
            f'<Relationship Id="rId{n+1}" Type="http://schemas.openxmlformats.org/'
            f'officeDocument/2006/relationships/styles" Target="styles.xml"/>')
        wb_rels.append("</Relationships>")

        styles = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                  '<fonts count="2"><font><sz val="11"/><name val="Calibri"/></font>'
                  f'<font><b/><sz val="11"/><color rgb="FF{HEADER_FONT}"/><name val="Calibri"/></font></fonts>'
                  '<fills count="3"><fill><patternFill patternType="none"/></fill>'
                  '<fill><patternFill patternType="gray125"/></fill>'
                  f'<fill><patternFill patternType="solid"><fgColor rgb="FF{HEADER_FILL}"/>'
                  '<bgColor indexed="64"/></patternFill></fill></fills>'
                  '<borders count="1"><border/></borders>'
                  '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
                  '<cellXfs count="2">'
                  '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
                  '<xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0" applyFont="1" applyFill="1"/>'
                  '</cellXfs></styleSheet>')

        root_rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                     '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                     '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
                     'officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
                     '</Relationships>')

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("[Content_Types].xml", "".join(content_types))
            z.writestr("_rels/.rels", root_rels)
            z.writestr("xl/workbook.xml", "".join(workbook))
            z.writestr("xl/_rels/workbook.xml.rels", "".join(wb_rels))
            z.writestr("xl/styles.xml", styles)
            for i, (_, rows) in enumerate(self.sheets, start=1):
                z.writestr(f"xl/worksheets/sheet{i}.xml", self._sheet_xml(rows))
        return buf.getvalue()


# ======================================================================================
# OPENPYXL BACKEND
# ======================================================================================
def _openpyxl_bytes(sheets) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    fill = PatternFill("solid", fgColor=HEADER_FILL)
    font = Font(bold=True, color=HEADER_FONT, size=10)

    for name, rows, widths in sheets:
        ws = wb.create_sheet(title=name[:31])
        if not rows:
            continue
        for c_i, value in enumerate(rows[0], start=1):
            cell = ws.cell(row=1, column=c_i, value=value)
            cell.fill, cell.font = fill, font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for r_i, row in enumerate(rows[1:], start=2):
            for c_i, value in enumerate(row, start=1):
                if value == "":
                    value = None
                ws.cell(row=r_i, column=c_i, value=value)
        for c_i in range(1, len(rows[0]) + 1):
            w = (widths[c_i - 1] if widths and c_i - 1 < len(widths) else None)
            if w is None:
                longest = max((len(str(rows[r][c_i - 1])) for r in range(min(len(rows), 200))
                               if c_i - 1 < len(rows[r]) and rows[r][c_i - 1] is not None), default=10)
                w = min(max(longest + 2, 10), 46)
            ws.column_dimensions[get_column_letter(c_i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ======================================================================================
# PUBLIC API
# ======================================================================================
def build_workbook(sheets) -> tuple[bytes, str]:
    """
    sheets: list of (name, headers, rows, widths_or_None)
    Returns (xlsx_bytes, backend_used).
    """
    try:
        payload = [(name, [headers] + rows, widths) for name, headers, rows, widths in sheets]
        return _openpyxl_bytes(payload), "openpyxl"
    except ImportError:
        wb = MiniXlsx()
        for name, headers, rows, _widths in sheets:
            wb.add_sheet(name, [headers] + rows)
        return wb.save_bytes(), "stdlib-fallback"
