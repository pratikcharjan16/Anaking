"""Walk-through question: the narrated, animated sequence a respondent watches.

A *walk-through* is a question that stores no answer.  It plays a sequence of steps -
each one an animated scene, or an image / video the author attached - with narration
spoken on top, and then the survey moves on.  It exists so a respondent can be *shown*
something instead of having to read it.

Two things are generated rather than uploaded:

* the **narration** - the author writes the text for a step and the respondent's browser
  speaks it (``speechSynthesis``) unless a recorded clip is attached instead.  Nothing is
  synthesised server-side, so a study works offline and needs no audio files.
* the **steps** - a PowerPoint deck can be turned into one step per slide, with the slide
  title as the step label and the slide body plus its speaker notes as the narration.

``steps_from_pptx`` is that second part.  It reads text only: images inside the deck are
left for the author to attach per step, because extracting them would mean re-hosting
files the Studio already has an upload path for.
"""

from __future__ import annotations

from typing import Any

# A deck bigger than this is almost certainly the wrong file, and reading it would eat the
# request budget.  PowerPoint decks are zip archives of XML, so the check is on size only.
DECK_MAX_BYTES = 12 * 1024 * 1024


class DeckError(ValueError):
    """A file that is not a readable PowerPoint deck."""


def _clean(text: str) -> str:
    """Collapse a slide's text into one narration line."""
    return " ".join((text or "").split())


def steps_from_pptx(data: bytes, max_steps: int = 60) -> list[dict[str, Any]]:
    """One walk-through step per slide in a ``.pptx`` deck.

    The slide's title becomes the step label; every other text shape on the slide, plus
    the speaker notes, becomes the narration the respondent hears.  Slides with no text at
    all are skipped rather than turned into a silent step nobody can follow.

    Returns a list of ``{label, narration}`` dicts in slide order, or raises
    :class:`DeckError` when the file cannot be read as a deck.
    """
    if not data:
        raise DeckError("the file is empty")
    if len(data) > DECK_MAX_BYTES:
        raise DeckError("the deck is too large (max %d MB)" % (DECK_MAX_BYTES // (1024 * 1024)))

    try:
        import io

        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - python-pptx ships with the app
        raise DeckError("PowerPoint reading is not installed on this server") from exc

    try:
        prs = Presentation(io.BytesIO(data))
        slides = list(prs.slides)
    except Exception as exc:  # zipfile.BadZipFile, ValueError, KeyError ...
        raise DeckError("that is not a readable PowerPoint deck") from exc

    if not slides:
        raise DeckError("the deck has no slides")

    steps: list[dict[str, Any]] = []
    for index, slide in enumerate(slides, start=1):
        title, body, notes = "", [], ""
        try:
            title = _clean(getattr(slide.shapes, "title", None) and
                           (slide.shapes.title.text_frame.text if slide.shapes.title.has_text_frame else "")
                           or "")
        except Exception:
            title = ""
        try:
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                # the title shape already became the label
                if title and shape == slide.shapes.title:
                    continue
                text = _clean(shape.text_frame.text)
                if text:
                    body.append(text)
        except Exception:
            pass
        try:
            if slide.has_notes_slide:
                notes = _clean(slide.notes_slide.notes_text_frame.text or "")
        except Exception:
            notes = ""

        narration = " ".join(body + ([notes] if notes else []))
        narration = _clean(narration)
        if not title and not narration:
            continue  # a blank slide would leave the respondent staring at nothing
        steps.append({
            "label": title or "Slide %d" % index,
            "narration": narration,
        })
        if len(steps) >= max_steps:
            break

    if not steps:
        raise DeckError("no slide in the deck has any text to narrate")
    return steps


def walk_steps(q: dict) -> list[dict]:
    """The steps of a walk-through question, in order, with the keys the player needs.

    Tolerates a question saved by an older build (``walkthrough_steps``) and fills in the
    defaults the explainer player expects so the client never has to guess.
    """
    raw = q.get("walk_steps")
    if not isinstance(raw, list):
        raw = q.get("walkthrough_steps")
    if not isinstance(raw, list):
        return []
    out = []
    for i, s in enumerate(raw):
        if not isinstance(s, dict):
            continue
        step = dict(s)
        step["label"] = str(step.get("label") or "Step %d" % (i + 1))
        # the explainer speaks `caption`; `title` is the on-screen heading
        step["title"] = str(step.get("title") or step["label"])
        step["caption"] = str(step.get("narration") or step.get("caption") or "")
        if not step.get("art"):
            step["art"] = "generic"
        out.append(step)
    return out


def walk_questions(q: dict) -> list[dict]:
    """The questions asked about a walk-through, once it has played."""
    raw = q.get("walk_questions")
    if not isinstance(raw, list):
        raw = q.get("walkthrough_questions")
    return [x for x in (raw or []) if isinstance(x, dict)]


def walk_scenes(q: dict, tts: bool | None = None) -> list[dict]:
    """Steps shaped for ``BeaconExplainer`` - what the respondent actually plays."""
    scenes = walk_steps(q)
    if tts is not None:
        for s in scenes:
            s["tts"] = bool(tts)
    return scenes
