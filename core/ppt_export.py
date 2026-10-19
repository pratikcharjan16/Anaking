"""
PowerPoint (.pptx) report generator for survey research findings.
Generates fully editable PowerPoint presentations featuring native editable charts,
editable data tables, KPI summary cards, and data-driven insights.
Supports full survey report, single question slide, or customized question sets.
"""

from __future__ import annotations

import io
from datetime import datetime
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData

# Brand palette
NAVY = RGBColor(15, 43, 72)        # #0f2b48
TEAL = RGBColor(18, 120, 158)      # #12789e
LIGHT_TEAL = RGBColor(238, 246, 249) # #eef6f9
AMBER = RGBColor(245, 158, 11)     # #f59e0b
EMERALD = RGBColor(16, 185, 129)   # #10b981
SLATE = RGBColor(30, 41, 59)       # #1e293b
MUTED = RGBColor(100, 116, 139)    # #64748b
LIGHT_GRAY = RGBColor(241, 245, 249) # #f1f5f9
WHITE = RGBColor(255, 255, 255)


def _set_cell_text(cell, text: str, bold: bool = False, size: int = 11,
                    color: RGBColor = SLATE, align: PP_ALIGN = PP_ALIGN.LEFT):
    cell.text = str(text)
    for p in cell.text_frame.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.name = "Segoe UI"
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color


def build_ppt_report(study_title: str, summary_data: dict, questions_to_include: list[str] | None = None,
                     scope: str = "real") -> bytes:
    """Builds a complete, professional, editable PowerPoint presentation for the survey.

    questions_to_include: optional list of question IDs to include. If None, includes all questions.
    """
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    all_qs = summary_data.get("questions") or []
    if questions_to_include:
        target_ids = {str(qid).strip().upper() for qid in questions_to_include}
        selected_qs = [q for q in all_qs if str(q.get("id")).strip().upper() in target_ids]
    else:
        selected_qs = all_qs

    cut_info = summary_data.get("data_cut")
    is_cut = cut_info and cut_info.get("active")

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide (only for full report or multi-question sets)
    # -------------------------------------------------------------
    is_single_q = len(selected_qs) == 1 and bool(questions_to_include)

    if not is_single_q:
        slide = prs.slides.add_slide(blank_layout)

        # Header card background
        top_bar = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = NAVY
        top_bar.line.fill.background()

        # Title text inside banner
        tb = slide.shapes.add_textbox(Inches(1.4), Inches(1.5), Inches(10.5), Inches(2.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p0 = tf.paragraphs[0]
        p0.text = "RESEARCH SURVEY REPORT"
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = AMBER

        p1 = tf.add_paragraph()
        p1.text = study_title
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(32)
        p1.font.bold = True
        p1.font.color.rgb = WHITE

        p2 = tf.add_paragraph()
        p2.text = "Executive Results & Dynamic Live Data Analysis"
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(16)
        p2.font.color.rgb = RGBColor(180, 215, 235)

        # Metadata badges on title slide
        meta_box = slide.shapes.add_textbox(Inches(1.4), Inches(4.3), Inches(10.5), Inches(1.5))
        mtf = meta_box.text_frame
        mtf.word_wrap = True

        now_str = datetime.now().strftime("%B %d, %Y")
        n_complete = summary_data.get("complete", 0)
        n_started = summary_data.get("started", 0)

        mp = mtf.paragraphs[0]
        mp.text = f"Date: {now_str}   |   Data Scope: {scope.capitalize()} Responses   |   Total Starts: {n_started}   |   Completed Sample: N = {n_complete}"
        mp.font.name = "Segoe UI"
        mp.font.size = Pt(13)
        mp.font.color.rgb = RGBColor(220, 235, 245)

        if is_cut:
            mp2 = mtf.add_paragraph()
            mp2.text = f"Data Cut Active: Cut by {cut_info['cut_q']} ({cut_info['segA_label']} n={cut_info['segA_n']} vs {cut_info['segB_label']} n={cut_info['segB_n']}) at {cut_info['conf_pct']}% Confidence Level"
            mp2.font.name = "Segoe UI"
            mp2.font.size = Pt(12)
            mp2.font.bold = True
            mp2.font.color.rgb = AMBER

    # -------------------------------------------------------------
    # SLIDE 2: Executive Summary & KPIs (if not single question)
    # -------------------------------------------------------------
    if not is_single_q:
        slide_kpi = prs.slides.add_slide(blank_layout)

        # Title
        t_box = slide_kpi.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.733), Inches(0.8))
        ktf = t_box.text_frame
        kp = ktf.paragraphs[0]
        kp.text = "Executive Summary & Fieldwork Metrics"
        kp.font.name = "Segoe UI"
        kp.font.size = Pt(24)
        kp.font.bold = True
        kp.font.color.rgb = NAVY

        # KPI Metric Cards
        kpi_metrics = [
            ("TOTAL STARTED", str(summary_data.get("started", 0)), f"{scope.capitalize()} scope"),
            ("COMPLETED", str(summary_data.get("complete", 0)), f"{summary_data.get('completion_pct', 0)}% completion rate"),
            ("MEDIAN DURATION", f"{summary_data.get('median_minutes', '—')} min", "Completed respondents"),
            ("SCREENED OUT", str(summary_data.get("screened", 0)), "Excluded from analysis"),
        ]

        card_w = Inches(2.7)
        card_h = Inches(1.8)
        start_x = Inches(0.8)
        gap = Inches(0.3)

        for i, (k_label, k_val, k_sub) in enumerate(kpi_metrics):
            cx = start_x + i * (card_w + gap)
            card = slide_kpi.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(1.5), card_w, card_h)
            card.fill.solid()
            card.fill.fore_color.rgb = LIGHT_TEAL
            card.line.color.rgb = TEAL

            ctb = slide_kpi.shapes.add_textbox(cx + Inches(0.15), Inches(1.65), card_w - Inches(0.3), card_h - Inches(0.3))
            ctf = ctb.text_frame
            ctf.word_wrap = True

            cp0 = ctf.paragraphs[0]
            cp0.text = k_label
            cp0.font.name = "Segoe UI"
            cp0.font.size = Pt(11)
            cp0.font.bold = True
            cp0.font.color.rgb = TEAL

            cp1 = ctf.add_paragraph()
            cp1.text = k_val
            cp1.font.name = "Segoe UI"
            cp1.font.size = Pt(28)
            cp1.font.bold = True
            cp1.font.color.rgb = NAVY

            cp2 = ctf.add_paragraph()
            cp2.text = k_sub
            cp2.font.name = "Segoe UI"
            cp2.font.size = Pt(10)
            cp2.font.color.rgb = MUTED

        # Overview Insights Box
        info_card = slide_kpi.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.7), Inches(11.733), Inches(3.1))
        info_card.fill.solid()
        info_card.fill.fore_color.rgb = LIGHT_GRAY
        info_card.line.fill.background()

        itb = slide_kpi.shapes.add_textbox(Inches(1.1), Inches(3.9), Inches(11.1), Inches(2.7))
        itf = itb.text_frame
        itf.word_wrap = True

        ip0 = itf.paragraphs[0]
        ip0.text = "Key Research Highlights & Methodology Note"
        ip0.font.name = "Segoe UI"
        ip0.font.size = Pt(16)
        ip0.font.bold = True
        ip0.font.color.rgb = NAVY

        ip1 = itf.add_paragraph()
        ip1.text = f"• Survey Instrument: {len(all_qs)} total questions evaluated across n={summary_data.get('complete', 0)} completed responses."
        ip1.font.name = "Segoe UI"
        ip1.font.size = Pt(13)
        ip1.font.color.rgb = SLATE

        ip2 = itf.add_paragraph()
        if is_cut:
            ip2.text = f"• Segment Comparison: Evaluated {cut_info['segA_label']} (n={cut_info['segA_n']}) versus {cut_info['segB_label']} (n={cut_info['segB_n']}) with two-proportion z-tests and Welch's t-tests."
        else:
            ip2.text = "• Descriptive Analytics: All answer bases reflect active respondents; missing/skipped items excluded from bases."
        ip2.font.name = "Segoe UI"
        ip2.font.size = Pt(13)
        ip2.font.color.rgb = SLATE

        ip3 = itf.add_paragraph()
        ip3.text = "• Native Editable Charts: Charts on following slides are fully native PowerPoint charts. Right-click any chart to edit data or colors in Excel."
        ip3.font.name = "Segoe UI"
        ip3.font.size = Pt(12)
        ip3.font.color.rgb = MUTED

    # -------------------------------------------------------------
    # QUESTION SLIDES
    # -------------------------------------------------------------
    for q in selected_qs:
        qid = q.get("id", "")
        qtitle = q.get("title", "")
        qtype = q.get("type", "")
        base_n = q.get("base", 0)
        insight_text = q.get("insight") or f"Answered base: {base_n} respondents."

        slide_q = prs.slides.add_slide(blank_layout)

        # Question Header Box
        h_box = slide_q.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.85))
        htf = h_box.text_frame
        htf.word_wrap = True

        hp0 = htf.paragraphs[0]
        hp0.text = f"{qid} · {qtitle}"
        hp0.font.name = "Segoe UI"
        hp0.font.size = Pt(20)
        hp0.font.bold = True
        hp0.font.color.rgb = NAVY

        hp1 = htf.add_paragraph()
        cut_label = f"   |   Data Cut: {cut_info['segA_label']} vs {cut_info['segB_label']}" if is_cut else ""
        hp1.text = f"Answered base: n = {base_n} respondents   |   Question Type: {qtype.replace('_', ' ').capitalize()}{cut_label}"
        hp1.font.name = "Segoe UI"
        hp1.font.size = Pt(11)
        hp1.font.color.rgb = MUTED

        # Insight Callout Banner
        ins_card = slide_q.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.75))
        ins_card.fill.solid()
        ins_card.fill.fore_color.rgb = LIGHT_TEAL
        ins_card.line.color.rgb = TEAL
        ins_card.line.width = Pt(1.5)

        ins_tb = slide_q.shapes.add_textbox(Inches(1.0), Inches(1.38), Inches(11.3), Inches(0.7))
        intf = ins_tb.text_frame
        intf.word_wrap = True
        inp = intf.paragraphs[0]
        inp.text = f"KEY INSIGHT:  {insight_text}"
        inp.font.name = "Segoe UI"
        inp.font.size = Pt(11.5)
        inp.font.bold = True
        inp.font.color.rgb = NAVY

        # Determine Chart Data and Table Data based on Question Type
        opts = q.get("options") or []
        rows = q.get("rows") or []

        has_chart = False
        chart_type = XL_CHART_TYPE.BAR_CLUSTERED
        cdata = CategoryChartData()

        table_headers = []
        table_rows = []

        # 1. Single Select & Multi Select
        if qtype in ("single_select", "multi_select") and opts:
            has_chart = True
            cdata.categories = [o["label"][:28] for o in opts]
            if is_cut:
                segA_vals = tuple((o.get("segA_pct") or 0) for o in opts)
                segB_vals = tuple((o.get("segB_pct") or 0) for o in opts)
                cdata.add_series(cut_info["segA_label"], segA_vals)
                cdata.add_series(cut_info["segB_label"], segB_vals)
                table_headers = ["Answer Option", f"{cut_info['segA_label']} (%)", f"{cut_info['segB_label']} (%)", "Diff (%)", "Sig"]
                for o in opts:
                    sig = "★ Sig" if (o.get("sig_test") and o["sig_test"].get("significant")) else "—"
                    diff = o.get("sig_test", {}).get("diff_pct", 0) if o.get("sig_test") else 0
                    table_rows.append([o["label"], f"{o.get('segA_pct', 0)}%", f"{o.get('segB_pct', 0)}%", f"{diff:+0.1f}%", sig])
            else:
                pct_vals = tuple((o.get("pct") or 0) for o in opts)
                cdata.add_series("Share (%)", pct_vals)
                table_headers = ["Answer Option", "Count (n)", "Share (%)"]
                for o in opts:
                    table_rows.append([o["label"], str(o.get("count", 0)), f"{o.get('pct', 0)}%"])

        # 2. Rating Scale, Grids, or Gauge with Rows
        elif (qtype in ("rating_scale", "rating_grid", "semantic_diff", "numeric_matrix") or (qtype == "gauge_meter" and rows)) and rows:
            has_chart = True
            cdata.categories = [r["label"][:28] for r in rows]
            if is_cut:
                segA_means = tuple((r.get("segA_mean") or 0) for r in rows)
                segB_means = tuple((r.get("segB_mean") or 0) for r in rows)
                cdata.add_series(cut_info["segA_label"], segA_means)
                cdata.add_series(cut_info["segB_label"], segB_means)
                table_headers = ["Item / Row", f"{cut_info['segA_label']} Mean", f"{cut_info['segB_label']} Mean", "Diff", "Sig"]
                for r in rows:
                    sig = "★ Sig" if (r.get("sig_test") and r["sig_test"].get("significant")) else "—"
                    diff = r.get("sig_test", {}).get("diff_mean", 0) if r.get("sig_test") else 0
                    table_rows.append([r["label"], str(r.get("segA_mean") or "—"), str(r.get("segB_mean") or "—"), f"{diff:+0.2f}", sig])
            else:
                means = tuple((r.get("mean") or 0) for r in rows)
                cdata.add_series("Mean Score", means)
                table_headers = ["Item / Row", "Base (n)", "Mean Score"]
                for r in rows:
                    table_rows.append([r["label"], str(r.get("base", 0)), f"{r.get('mean', 0):.2f}" if r.get("mean") is not None else "—"])

        # 3. Single Gauge Meter
        elif qtype == "gauge_meter" and not rows:
            has_chart = True
            zones = q.get("zones") or []
            if zones:
                cdata.categories = [z["label"] for z in zones]
                cdata.add_series("Respondents (%)", tuple(z["pct"] for z in zones))
            table_headers = ["Metric / Zone", "Value"]
            table_rows.append(["Average Gauge Score", f"{q.get('mean', '—')} {q.get('unit', '%')}"])
            table_rows.append(["Target Benchmark", f"{q.get('target', '—')} {q.get('unit', '%')}"])
            for z in zones:
                table_rows.append([f"Zone: {z['label']}", f"{z['pct']}% (n={z['count']})"])

        # 4. Numeric / Slider / NPS
        elif qtype in ("numeric", "slider", "nps"):
            table_headers = ["Metric", "Value"]
            table_rows.append(["Answered Base", str(q.get("base", 0))])
            table_rows.append(["Mean", str(q.get("mean", "—"))])
            table_rows.append(["Median", str(q.get("median", "—"))])
            if q.get("std_dev") is not None:
                table_rows.append(["Std Deviation", str(q["std_dev"])])
            if q.get("min_val") is not None and q.get("max_val") is not None:
                table_rows.append(["Min / Max Range", f"{q['min_val']} - {q['max_val']}"])

        # Render Chart on Left if available
        if has_chart:
            chart_shape = slide_q.shapes.add_chart(
                chart_type, Inches(0.8), Inches(2.3), Inches(6.4), Inches(4.7), cdata
            )
            ch = chart_shape.chart
            ch.has_legend = True
            ch.legend.position = XL_LEGEND_POSITION.TOP
            ch.has_title = True
            ch.chart_title.text_frame.text = "Distribution / Scores"
            ch.plots[0].has_data_labels = True

            # Table on Right
            tbl_x = Inches(7.5)
            tbl_w = Inches(5.0)
        else:
            # Full width Table
            tbl_x = Inches(0.8)
            tbl_w = Inches(11.733)

        if table_headers and table_rows:
            t_rows_n = min(12, len(table_rows) + 1)
            table_shape = slide_q.shapes.add_table(t_rows_n, len(table_headers), tbl_x, Inches(2.3), tbl_w, Inches(4.7))
            tbl = table_shape.table

            for col_i, h in enumerate(table_headers):
                _set_cell_text(tbl.cell(0, col_i), h, bold=True, size=10, color=WHITE, align=PP_ALIGN.CENTER if col_i > 0 else PP_ALIGN.LEFT)
                tbl.cell(0, col_i).fill.solid()
                tbl.cell(0, col_i).fill.fore_color.rgb = NAVY

            for row_i, rdata in enumerate(table_rows[:t_rows_n - 1]):
                bg_col = WHITE if row_i % 2 == 0 else LIGHT_GRAY
                for col_i, val in enumerate(rdata):
                    cell = tbl.cell(row_i + 1, col_i)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = bg_col
                    is_sig = "★" in str(val)
                    _set_cell_text(cell, val, bold=is_sig, size=10, color=AMBER if is_sig else SLATE, align=PP_ALIGN.CENTER if col_i > 0 else PP_ALIGN.LEFT)

    bio = io.BytesIO()
    prs.save(bio)
    bio.seek(0)
    return bio.getvalue()
